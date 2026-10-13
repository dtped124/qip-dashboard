"""
Permission layer tests (Phase 2: viewer role + backend permissions).

Guard rails
  test_every_api_route_refuses_anonymous    no /api/ route is open to anonymous callers
                                            (only /api/auth/login and /api/auth/options are exempt)
  test_matrix_declares_every_api_route      every route has a decision in MATRIX, so a new
                                            endpoint cannot ship without choosing a role
  test_matrix_covers_every_drf_method       ... and so does every HTTP method a DRF view accepts
  test_every_write_route_enforces_csrf      every write route is CSRF-protected

Decisions
  test_matrix                               each (method, route) x account answers as declared

Everything else covers the special cases (staff flag, superuser, deactivated accounts,
stale roles, error bodies, CSRF, patient-number masking, the entry-benchmarks fix).

"Allowed" means anything but 401/403 and no server error: 400/404/405/415 are fine (the view
ran and rejected the empty request), and so is the deliberate 501 of the HIS placeholders.
"Denied" means exactly 403 with code permission_denied.
"""
from __future__ import annotations

import datetime
import io
import json
import re

import pytest
from django.contrib.auth.models import AnonymousUser
from django.http import HttpResponse
from django.urls import URLResolver, get_resolver
from django.urls.resolvers import RoutePattern
from openpyxl import load_workbook

from apps.accounts.models import User, UserRole
from apps.accounts.permissions import has_role, is_admin, require_roles
from apps.davinci.models import DavinciCase
from apps.entry.models import (
    Campus,
    CaseRecord,
    IndicatorAssignment,
    IndicatorEntry,
    MonthlyReport,
    ReportCategory,
)
from apps.indicators.models import Indicator, TCPIBenchmark

pytestmark = pytest.mark.django_db

# ── Route discovery ──────────────────────────────────────────────

PUBLIC_ROUTES = {"/api/auth/login", "/api/auth/options"}

_PARAM = re.compile(r"<(?:(?P<conv>\w+):)?(?P<name>\w+)>")
_FILL_BY_NAME = {"code": "HA01-01", "campus": "zhubei"}
_FILL_BY_CONVERTER = {"int": "999999", "uuid": "00000000-0000-0000-0000-000000000000"}


def fill(route: str) -> str:
    """Concrete URL for a route template. Ids are ones that cannot exist, so a
    request that gets past the permission check can never touch real rows."""

    def repl(m):
        return _FILL_BY_NAME.get(m["name"]) or _FILL_BY_CONVERTER.get(m["conv"], "x")

    return _PARAM.sub(repl, route)


def api_routes() -> list[tuple[str, object]]:
    """(route template, view callback) for every URL under /api/, e.g.
    ("/api/admin/users/<int:pk>", <view>). The Django admin site is skipped."""
    found: list[tuple[str, object]] = []

    def walk(patterns, prefix: str) -> None:
        for entry in patterns:
            route = prefix + str(entry.pattern)
            if isinstance(entry, URLResolver):
                if route.startswith("/api/") or "/api/".startswith(route):
                    walk(entry.url_patterns, route)
            elif route.startswith("/api/"):
                assert isinstance(entry.pattern, RoutePattern), f"regex route not supported by this test: {route}"
                found.append((route, entry.callback))

    walk(get_resolver().url_patterns, "/")
    return found


def call(client, method: str, url: str):
    if method == "GET":
        return client.get(url)
    return client.generic(method, url, data="{}", content_type="application/json")


def let_in(resp) -> bool:
    """The view was reached: not an auth rejection, and not a crash (501 = intended placeholder)."""
    return resp.status_code not in (401, 403) and not (500 <= resp.status_code < 600 and resp.status_code != 501)


def body(resp) -> dict:
    try:
        data = resp.json()
    except ValueError:
        return {}
    return data if isinstance(data, dict) else {}


# ── The permission matrix ────────────────────────────────────────
# account names: viewer, reporter, reviewer, admin, norole (signed in, no roles)

READERS = frozenset({"viewer", "reviewer", "admin"})
ADMIN_ONLY = frozenset({"admin"})
REPORTERS = frozenset({"reporter", "admin"})
REVIEWERS = frozenset({"reviewer", "admin"})
REPORTER_OR_REVIEWER = frozenset({"reporter", "reviewer", "admin"})
ANY_SIGNED_IN = frozenset({"viewer", "reporter", "reviewer", "admin", "norole"})

# (method, route template as in urls.py, accounts that are let in)
MATRIX: list[tuple[str, str, frozenset[str]]] = [
    # auth (login itself is public: see PUBLIC_ROUTES)
    ("GET", "/api/auth/me", ANY_SIGNED_IN),
    ("POST", "/api/auth/logout", ANY_SIGNED_IN),
    ("POST", "/api/auth/change-password", ANY_SIGNED_IN),
    # QIP dashboards and indicator data: readers; writes and full export: admin
    ("GET", "/api/v1/dashboard/", READERS),
    ("GET", "/api/v1/indicators/", READERS),
    ("GET", "/api/v1/indicators/export/", ADMIN_ONLY),
    ("GET", "/api/v1/indicators/export/element-list/", READERS),
    ("GET", "/api/v1/indicators/<str:code>/", READERS),
    ("PATCH", "/api/v1/indicators/<str:code>/", ADMIN_ONLY),
    ("GET", "/api/v1/indicators/<str:code>/data/", READERS),
    ("GET", "/api/v1/indicators/<str:code>/alerts/", READERS),
    ("GET", "/api/v1/indicators/<str:code>/summaries/", READERS),
    ("GET", "/api/v1/indicators/<str:code>/analysis/", READERS),
    ("GET", "/api/v1/tcpi/", READERS),
    ("POST", "/api/v1/tcpi/", ADMIN_ONLY),
    # Excel imports of the dashboard data: admin
    ("POST", "/api/v1/imports/upload/", ADMIN_ONLY),
    ("GET", "/api/v1/imports/logs/", ADMIN_ONLY),
    ("POST", "/api/v1/imports/correct-datapoint/", ADMIN_ONLY),
    # dashboard data from the entry system: readers
    ("GET", "/api/dashboard/entry-data", READERS),
    ("GET", "/api/dashboard/entry-benchmarks", READERS),
    # administration: admin
    ("GET", "/api/admin/users", ADMIN_ONLY),
    ("POST", "/api/admin/users", ADMIN_ONLY),
    ("GET", "/api/admin/users/<int:pk>", ADMIN_ONLY),
    ("PATCH", "/api/admin/users/<int:pk>", ADMIN_ONLY),
    ("DELETE", "/api/admin/users/<int:pk>", ADMIN_ONLY),
    ("POST", "/api/admin/users/<int:pk>/reset-password", ADMIN_ONLY),
    ("GET", "/api/admin/campuses", ADMIN_ONLY),
    ("GET", "/api/admin/categories", ADMIN_ONLY),
    ("GET", "/api/admin/assignments", ADMIN_ONLY),
    ("POST", "/api/admin/assignments", ADMIN_ONLY),
    ("GET", "/api/admin/assignments/<int:pk>", ADMIN_ONLY),
    ("PATCH", "/api/admin/assignments/<int:pk>", ADMIN_ONLY),
    ("DELETE", "/api/admin/assignments/<int:pk>", ADMIN_ONLY),
    ("GET", "/api/admin/deadlines", ADMIN_ONLY),
    ("POST", "/api/admin/deadlines", ADMIN_ONLY),
    # data entry: reporter
    ("GET", "/api/entry/my-tasks", REPORTERS),
    ("GET", "/api/entry/form", REPORTERS),
    ("POST", "/api/entry/save-draft", REPORTERS),
    ("POST", "/api/entry/submit", REPORTERS),
    # review: reviewer (unlock: admin)
    ("GET", "/api/review/overview", REVIEWERS),
    ("GET", "/api/review/detail", REVIEWERS),
    ("POST", "/api/review/approve", REVIEWERS),
    ("POST", "/api/review/reject", REVIEWERS),
    ("PATCH", "/api/review/edit-entry", REVIEWERS),
    ("POST", "/api/review/finalize", REVIEWERS),
    ("POST", "/api/review/unlock", ADMIN_ONLY),
    # entry-system imports: reviewer
    ("POST", "/api/import/excel", REVIEWERS),
    ("POST", "/api/import/confirm", REVIEWERS),
    ("GET", "/api/import/batches", REVIEWERS),
    ("POST", "/api/import/his-trigger", REVIEWERS),
    ("POST", "/api/import/his-webhook", REVIEWERS),
    # case lists: reporter works them, reviewer may read and review them
    ("GET", "/api/case-list/", REPORTER_OR_REVIEWER),
    ("POST", "/api/case-list/exclude", REPORTERS),
    ("POST", "/api/case-list/restore", REPORTERS),
    ("GET", "/api/case-list/exclusion-reasons", REPORTER_OR_REVIEWER),
    ("POST", "/api/case-list/review-exclusion", REVIEWERS),
    # davinci: readers see aggregates, reviewer|admin see case-level rows, admin imports
    ("POST", "/api/davinci/import/", ADMIN_ONLY),
    ("POST", "/api/davinci/import/confirm/", ADMIN_ONLY),
    ("GET", "/api/davinci/import/logs/", ADMIN_ONLY),
    ("GET", "/api/davinci/indicators/", READERS),
    ("GET", "/api/davinci/indicators/<str:code>/series/", READERS),
    ("GET", "/api/davinci/drilldown/", READERS),
    ("GET", "/api/davinci/cases/", REVIEWERS),
    ("GET", "/api/davinci/export/", REVIEWERS),
    ("GET", "/api/davinci/meta/", READERS),
]


def _matrix_id(row) -> str:
    return f"{row[0]} {row[1]}"


# ── Guard rails ──────────────────────────────────────────────────


def test_every_api_route_refuses_anonymous(client_for):
    """Anonymous callers get 401/403 from every /api/ route, whatever the method.

    (Stricter than "GET, or POST for POST-only routes": all write methods are tried
    too, so a route that only guards GET still fails here.)
    """
    routes = [route for route, _ in api_routes() if route not in PUBLIC_ROUTES]
    assert len(routes) >= 50, f"route discovery looks broken: only {len(routes)} routes found"

    anon = client_for()
    failures = []
    for route in routes:
        url = fill(route)
        for method in ("GET", "POST", "PUT", "PATCH", "DELETE"):
            resp = call(anon, method, url)
            if resp.status_code not in (401, 403) or body(resp).get("code") != "not_authenticated":
                failures.append(f"{method} {url} -> {resp.status_code} {resp.content[:80]!r}")
    assert not failures, "routes reachable without signing in:\n" + "\n".join(failures)


def test_login_and_options_are_the_only_public_routes(client_for):
    """The exceptions to the meta-test really are reachable anonymously."""
    assert {r for r, _ in api_routes() if r in PUBLIC_ROUTES} == PUBLIC_ROUTES
    resp = call(client_for(), "POST", "/api/auth/login")
    assert resp.status_code == 400  # validation error, not 401/403
    assert "employee_id" in body(resp)
    options = call(client_for(), "GET", "/api/auth/options")
    assert options.status_code == 200
    assert body(options) == {"ldap_enabled": False}


def test_matrix_declares_every_api_route():
    discovered = {route for route, _ in api_routes()}
    declared = {route for _, route, _ in MATRIX} | PUBLIC_ROUTES
    assert discovered - declared == set(), "routes without a permission decision in MATRIX"
    assert declared - discovered == set(), "MATRIX rows for routes that no longer exist"


def test_matrix_covers_every_drf_method():
    """A DRF view that starts accepting another method must get a MATRIX row."""
    declared = {(method, route) for method, route, _ in MATRIX}
    missing = []
    for route, callback in api_routes():
        view_cls = getattr(callback, "cls", None)  # set by APIView.as_view()
        if view_cls is None or route in PUBLIC_ROUTES:
            continue
        for method in view_cls.http_method_names:
            if method not in ("options", "head") and (method.upper(), route) not in declared:
                missing.append(f"{method.upper()} {route}")
    assert not missing, "DRF methods without a MATRIX row: " + ", ".join(missing)


@pytest.mark.parametrize(
    "method,route",
    [(m, r) for m, r, _ in MATRIX if m != "GET"],
    ids=[f"{m} {r}" for m, r, _ in MATRIX if m != "GET"],
)
def test_every_write_route_enforces_csrf(method, route, admin_account, client_for):
    """Signed in as admin (so permissions cannot be the reason), a write without
    the CSRF token is rejected - for DRF and plain Django views alike."""
    client = client_for(admin_account, enforce_csrf_checks=True)
    resp = call(client, method, fill(route))
    assert resp.status_code == 403
    assert body(resp) == {"detail": "CSRF verification failed.", "code": "csrf_failed"}


# ── The matrix ───────────────────────────────────────────────────


@pytest.mark.parametrize("method,route,allowed", MATRIX, ids=[_matrix_id(r) for r in MATRIX])
def test_matrix(method, route, allowed, accounts, client_for):
    url = fill(route)
    problems = []
    for name, user in accounts.items():
        resp = call(client_for(user), method, url)
        if name in allowed:
            ok = let_in(resp)
            expected = "allowed"
        else:
            ok = resp.status_code == 403 and body(resp).get("code") == "permission_denied"
            expected = "403 permission_denied"
        if not ok:
            problems.append(f"  {name}: got {resp.status_code} {resp.content[:80]!r}, expected {expected}")
    assert not problems, f"{method} {url}\n" + "\n".join(problems)


# ── Special cases ────────────────────────────────────────────────


def test_staff_flag_is_not_admin(make_user, client_for):
    """is_staff only opens the Django admin site; it must grant nothing in the API."""
    staff = make_user("t_staff", [], is_staff=True)
    staff_reporter = make_user("t_staff_rep", ["reporter"], is_staff=True)
    for user in (staff, staff_reporter):
        client = client_for(user)
        assert client.get("/api/admin/users").status_code == 403
        assert client.post("/api/review/unlock", data="{}", content_type="application/json").status_code == 403
        assert client.get("/api/v1/dashboard/").status_code == 403
    assert client_for(staff).get("/api/auth/me").status_code == 200


def test_superuser_without_roles_counts_as_admin(make_user, client_for):
    superuser = make_user("t_super", [], is_superuser=True)
    client = client_for(superuser)
    assert client.get("/api/admin/users").status_code == 200
    assert client.get("/api/v1/dashboard/").status_code == 200


def test_deactivated_account_loses_access_immediately(admin_account, client_for):
    client = client_for(admin_account)
    assert client.get("/api/admin/users").status_code == 200
    User.objects.filter(pk=admin_account.pk).update(is_active=False)
    resp = client.get("/api/admin/users")
    assert resp.status_code == 403
    assert body(resp)["code"] == "not_authenticated"


def test_role_changes_apply_to_the_next_request(viewer, client_for):
    """Roles are read from the database on every request, not cached in the session."""
    client = client_for(viewer)
    assert client.get("/api/v1/dashboard/").status_code == 200
    User.objects.filter(pk=viewer.pk).update(roles=["reporter"])
    assert client.get("/api/v1/dashboard/").status_code == 403
    assert client.get("/api/entry/my-tasks").status_code != 403


def test_viewer_cannot_reach_case_lists(viewer, client_for):
    """Case lists carry patient numbers: viewers are readers of aggregates only."""
    client = client_for(viewer)
    assert client.get("/api/case-list/").status_code == 403
    assert client.get("/api/case-list/exclusion-reasons").status_code == 403


# Davinci: the cases list and the xlsx export (2nd sheet) are row-level data - account
# number, surgeon name - so they are reviewer | admin. Aggregates stay open to viewers.
DAVINCI_CASES = "/api/davinci/cases/?campus=竹北&period=202601"
DAVINCI_EXPORT = "/api/davinci/export/?campus=竹北"
DAVINCI_ROW_LEVEL = pytest.mark.parametrize("url", [DAVINCI_CASES, DAVINCI_EXPORT], ids=["cases", "export"])


@pytest.fixture
def davinci_case():
    return DavinciCase.objects.create(
        campus="竹北", period=202601, account="ACC-0001", dept_name="婦產科", surgeon="王醫師",
        op_time_min=100.0, blood_ml=10.0,
    )


@DAVINCI_ROW_LEVEL
def test_viewer_is_refused_davinci_case_level_data(url, davinci_case, viewer, client_for):
    resp = client_for(viewer).get(url)
    assert resp.status_code == 403
    assert body(resp)["code"] == "permission_denied"
    assert b"ACC-0001" not in resp.content


@DAVINCI_ROW_LEVEL
@pytest.mark.parametrize("who", ["reviewer", "admin"])
def test_reviewer_and_admin_are_let_in_to_davinci_case_level_data(url, who, davinci_case, accounts, client_for):
    resp = client_for(accounts[who]).get(url)
    assert let_in(resp)
    assert resp.status_code == 200
    if url == DAVINCI_CASES:
        assert [c["account"] for c in resp.json()["data"]] == ["ACC-0001"]
    else:
        sheet = load_workbook(io.BytesIO(resp.content))["個案明細"]
        assert "ACC-0001" in {cell.value for row in sheet.iter_rows() for cell in row}


def test_viewer_still_sees_davinci_aggregates(davinci_case, viewer, client_for):
    client = client_for(viewer)
    for url in (
        "/api/davinci/meta/",
        "/api/davinci/indicators/?campus=竹北",
        "/api/davinci/indicators/DV04/series/?campus=竹北",
    ):
        assert client.get(url).status_code == 200, url
    drilldown = client.get("/api/davinci/drilldown/?code=DV04&campus=竹北&period=202601&by=dept")
    assert drilldown.status_code == 200
    assert [(r["key"], r["denominator"]) for r in drilldown.json()["data"]] == [("婦產科", 1)]


class TestRoleHelpers:
    def test_anonymous_has_no_role(self):
        assert not is_admin(AnonymousUser())
        assert not has_role(AnonymousUser(), "viewer")
        assert not has_role(None, "viewer")

    def test_admin_role_passes_every_check(self):
        admin = User(employee_id="a", roles=["admin"])
        assert is_admin(admin)
        assert has_role(admin, "reporter")
        assert has_role(admin, "viewer", "reviewer")

    def test_other_roles_only_pass_their_own(self):
        viewer = User(employee_id="v", roles=["viewer"])
        assert has_role(viewer, "viewer")
        assert has_role(viewer, "reviewer", "viewer")
        assert not has_role(viewer, "reporter")
        assert not has_role(viewer, "admin")
        assert not is_admin(viewer)

    def test_superuser_counts_as_admin_but_staff_does_not(self):
        assert is_admin(User(employee_id="s", roles=[], is_superuser=True))
        assert not is_admin(User(employee_id="t", roles=[], is_staff=True))
        assert not has_role(User(employee_id="t", roles=[], is_staff=True), "viewer")

    def test_inactive_user_has_no_role(self):
        assert not has_role(User(employee_id="i", roles=["admin"], is_active=False), "admin")

    def test_enum_members_are_accepted(self):
        assert has_role(User(employee_id="v", roles=["viewer"]), UserRole.VIEWER)

    def test_malformed_roles_never_match(self):
        # a bare string would match by substring ("admin" in "xadminx")
        assert not is_admin(User(employee_id="m", roles="xadminx"))
        assert not has_role(User(employee_id="m", roles="viewer"), "viewer")
        assert not has_role(User(employee_id="m", roles={"admin": True}), "admin")
        assert not has_role(User(employee_id="m", roles=[None, 7]), "viewer")


def test_require_roles_without_roles_means_any_signed_in_user(rf):
    @require_roles()
    def view(request):
        return HttpResponse("ok")

    request = rf.get("/x")
    request.user = AnonymousUser()
    resp = view(request)
    assert resp.status_code == 403
    assert json.loads(resp.content) == {
        "detail": "Authentication credentials were not provided.",
        "code": "not_authenticated",
    }

    # A bare User() carries the model default must_change_password=True, which the gate refuses.
    request.user = User(employee_id="n", roles=[], must_change_password=False)
    assert view(request).status_code == 200


def test_require_roles_denies_wrong_role_with_permission_denied(rf):
    @require_roles("reviewer")
    def view(request):
        return HttpResponse("ok")

    request = rf.get("/x")
    request.user = User(employee_id="v", roles=["viewer"], must_change_password=False)
    resp = view(request)
    assert resp.status_code == 403
    assert json.loads(resp.content) == {
        "detail": "You do not have permission to perform this action.",
        "code": "permission_denied",
    }
    request.user = User(employee_id="r", roles=["reviewer"], must_change_password=False)
    assert view(request).status_code == 200
    request.user = User(employee_id="a", roles=["admin"], must_change_password=False)
    assert view(request).status_code == 200


# ── Error bodies ─────────────────────────────────────────────────


def test_error_bodies_carry_a_code(accounts, client_for):
    anon = client_for()

    drf_anon = anon.get("/api/v1/dashboard/")
    assert drf_anon.status_code == 403
    assert body(drf_anon)["code"] == "not_authenticated"
    assert "detail" in body(drf_anon)

    drf_denied = client_for(accounts["reporter"]).get("/api/v1/dashboard/")
    assert drf_denied.status_code == 403
    assert body(drf_denied)["code"] == "permission_denied"

    # plain Django views use the exact same shape
    plain_anon = anon.get("/api/davinci/meta/")
    assert plain_anon.status_code == 403
    assert body(plain_anon) == {
        "detail": "Authentication credentials were not provided.",
        "code": "not_authenticated",
    }

    # any other DRF exception with a single code
    not_allowed = client_for(accounts["norole"]).delete("/api/auth/me")
    assert not_allowed.status_code == 405
    assert body(not_allowed)["code"] == "method_not_allowed"

    # field-level validation errors keep their shape (no "code" added)
    invalid = call(anon, "POST", "/api/auth/login")
    assert invalid.status_code == 400
    assert "code" not in body(invalid)


# ── CSRF ─────────────────────────────────────────────────────────


def test_csrf_failure_on_a_drf_view_is_json(admin_account, client_for):
    client = client_for(admin_account, enforce_csrf_checks=True)
    resp = client.post("/api/v1/tcpi/", data="{}", content_type="application/json")
    assert resp.status_code == 403
    assert resp["Content-Type"].startswith("application/json")
    assert body(resp) == {"detail": "CSRF verification failed.", "code": "csrf_failed"}


def test_csrf_failure_on_a_plain_django_view_is_json(admin_account, client_for):
    client = client_for(admin_account, enforce_csrf_checks=True)
    resp = client.post("/api/davinci/import/")
    assert resp.status_code == 403
    assert resp["Content-Type"].startswith("application/json")
    assert body(resp) == {"detail": "CSRF verification failed.", "code": "csrf_failed"}


def test_requests_with_the_csrf_token_pass_both_view_kinds(admin_account, client_for):
    client = client_for(admin_account, enforce_csrf_checks=True)
    assert client.get("/api/auth/me").status_code == 200  # sets the csrftoken cookie
    token = client.cookies["csrftoken"].value

    drf = client.post(
        "/api/v1/tcpi/", data="{}", content_type="application/json", headers={"X-CSRFToken": token}
    )
    assert drf.status_code == 400  # reached the view: "No benchmarks provided"
    assert body(drf)["error"]["code"] == "BAD_REQUEST"

    plain = client.post("/api/davinci/import/", headers={"X-CSRFToken": token})
    assert plain.status_code == 400  # reached the view: "No file provided"
    assert body(plain)["error"]["code"] == "BAD_REQUEST"


def test_login_works_without_a_csrf_token(make_user, client_for, settings):
    settings.PASSWORD_HASHERS = ["django.contrib.auth.hashers.MD5PasswordHasher"]  # fast hashing for tests
    password = "test-only-pw-123"
    make_user("t_login", ["viewer"], password=password)

    client = client_for(enforce_csrf_checks=True)
    resp = client.post(
        "/api/auth/login",
        data=json.dumps({"employee_id": "t_login", "password": password}),
        content_type="application/json",
    )
    assert resp.status_code == 200
    assert body(resp)["roles"] == ["viewer"]
    assert client.get("/api/auth/me").status_code == 200  # the new session is usable


# ── Behaviour fixed together with the permission layer ───────────


def test_entry_benchmarks_serves_the_tcpi_payload(viewer, client_for):
    """Used to 500: a DRF Request was handed to another @api_view function."""
    indicator = Indicator.objects.create(
        code="HA01-01", name="t", category="整體照護", unit="percent", direction="lower", campuses=["竹北"]
    )
    TCPIBenchmark.objects.create(indicator=indicator, tcpi_name="t", year=113, regional_hospital=1.5)

    client = client_for(viewer)
    via_entry = client.get("/api/dashboard/entry-benchmarks")
    via_tcpi = client.get("/api/v1/tcpi/")
    assert via_entry.status_code == 200
    assert via_entry.json() == via_tcpi.json()
    assert via_entry.json()["total"] == 1


class TestCaseListPatientNumberMasking:
    """Who reads a case list, and who sees the full chart numbers in it.

    Reviewers and admins read every case list. A reporter reads the ones of indicators assigned to them
    (any field) at that campus in that month, and then sees the chart numbers in full; everybody else is
    refused outright. The account's own ``campus`` column and the Django is_staff flag decide nothing.
    (Detailed cases: apps/entry/tests/test_submit_and_review.py::TestCaseListReads.)
    """

    @pytest.fixture
    def url(self):
        zhubei = Campus.objects.create(code="zhubei", name="竹北", benchmark_level="regional")
        category = ReportCategory.objects.create(code="HA03", name="x", sort_order=1)
        report = MonthlyReport.objects.create(campus=zhubei, year=115, month=3, category=category)
        entry = IndicatorEntry.objects.create(report=report, indicator_code="HA03-01")
        CaseRecord.objects.create(entry=entry, case_role="numerator", his_raw_data={"chart_no": "A123456789"})
        return "/api/case-list/?indicator=HA03-01&campus=zhubei&year=115&month=3"

    def assign(self, user, code="HA03-01", campus="zhubei"):
        return IndicatorAssignment.objects.create(
            indicator_code=code, campus=Campus.objects.get(code=campus), user=user,
            effective_from=datetime.date(2026, 1, 1),
        )

    def chart_no(self, client, url) -> str:
        resp = client.get(url)
        assert resp.status_code == 200, resp.content
        return resp.json()["records"][0]["his_raw_data"]["chart_no"]

    def test_a_reporter_without_an_assignment_is_refused_even_on_the_same_campus(self, url, make_user, client_for):
        resident = make_user("t_rep_resident", ["reporter"], campus=Campus.objects.get(code="zhubei"))
        resp = client_for(resident).get(url)
        assert resp.status_code == 403
        assert body(resp)["code"] == "not_assigned"
        assert b"A123456789" not in resp.content

    def test_staff_flag_does_not_open_it(self, url, make_user, client_for):
        staff_reporter = make_user("t_rep_staff", ["reporter"], is_staff=True)
        resp = client_for(staff_reporter).get(url)
        assert resp.status_code == 403
        assert body(resp)["code"] == "not_assigned"

    def test_a_reporter_assigned_at_the_campus_sees_it_in_full(self, url, make_user, client_for):
        hsinchu = Campus.objects.create(code="hsinchu", name="新竹", benchmark_level="medical_center")
        own = make_user("t_rep_own", ["reporter"], campus=hsinchu)  # the account's campus is somewhere else
        self.assign(own)
        assert self.chart_no(client_for(own), url) == "A123456789"

    def test_an_assignment_at_another_campus_does_not_help(self, url, make_user, client_for):
        Campus.objects.create(code="hsinchu", name="新竹", benchmark_level="medical_center")
        other = make_user("t_rep_other", ["reporter"], campus=Campus.objects.get(code="zhubei"))
        self.assign(other, campus="hsinchu")
        assert client_for(other).get(url).status_code == 403

    def test_reviewer_and_admin_see_it_in_full(self, url, reviewer, admin_account, client_for):
        assert self.chart_no(client_for(reviewer), url) == "A123456789"
        assert self.chart_no(client_for(admin_account), url) == "A123456789"
